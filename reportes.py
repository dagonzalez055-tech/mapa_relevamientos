"""Reportes de los relevamientos: planilla de Excel y nota en Word.

Usa openpyxl y python-docx (se instalan con requirements.txt). No depende de Streamlit.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

import logica as lg

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre",
         "octubre", "noviembre", "diciembre"]
FUENTE_RED = "Extracción regulatoria de EDESA (INSTALACIONES_2025), red de media tensión de 13,2, 33 y 66 kV"


# ---------------------------------------------------------------- formato
def num(x: float, dec: int = 1) -> str:
    """Número con formato argentino (1.234,5)."""
    return f"{x:,.{dec}f}".replace(",", "\0").replace(".", ",").replace("\0", ".")


def fecha_larga(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def fecha_corta(texto) -> str:
    """'2026-09-03' → '03/09/2026'; vacío si no hay fecha."""
    try:
        a, m, d = str(texto).split("-")
        return f"{int(d):02d}/{int(m):02d}/{a}"
    except ValueError:
        return ""


def a_fecha(texto):
    """'2026-09-03' → date; None si no hay fecha."""
    try:
        return datetime.strptime(str(texto)[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# ------------------------------------------------------------------ datos
def _km(red: dict, segs, tension=None) -> float:
    segs = np.asarray(segs, dtype=np.int64)
    if len(segs) == 0:
        return 0.0
    m = np.ones(len(segs), dtype=bool) if tension is None else (red["tension"][segs] == np.float32(tension))
    return float(red["largo_m"][segs][m].sum()) / 1000


def datos_reporte(red: dict, rels: pd.DataFrame, tramos: dict, obs: pd.DataFrame, universo: pd.DataFrame,
                  fecha: date) -> dict:
    """Junta todo lo que van a mostrar el Excel y la nota.

    rels: relevamientos a incluir. tramos: {expediente: array de tramos de la red}. obs: observaciones.
    universo: lista de localidades de la provincia."""
    rels = rels.reset_index(drop=True)
    tabla_loc, fuera = lg.estado_localidades(universo, rels)
    depto = {}
    for _, u in universo.iterrows():
        for k in {lg.norm_nombre(u["nombre"])} | {lg.norm_nombre(a) for a in str(u["alias"]).split(";") if a.strip()}:
            depto[k] = u["departamento"]

    filas = []
    for _, r in rels.iterrows():
        e = r["expediente"]
        segs = np.asarray(tramos.get(e, []), dtype=np.int64)
        zona = ""
        if len(segs):
            z = np.bincount(red["area"][segs].astype(np.int64)).argmax()
            zona = str(red["areas"][z])
        deptos = []
        for parte in lg.partir_localidades(r["localidad"]):
            d = depto.get(lg.norm_nombre(parte), "")
            if d and d not in deptos:
                deptos.append(d)
        filas.append({
            "expediente": e, "localidad": r["localidad"], "departamento": ", ".join(deptos), "zona": zona,
            "desde": a_fecha(r["fecha_desde"]), "hasta": a_fecha(r["fecha_hasta"]), "informe": r["informe"],
            "modo": r["modo"], "km13": _km(red, segs, 13.2), "km33": _km(red, segs, 33.0), "km66": _km(red, segs, 66.0),
            "estado": r["estado"], "fecha_ejecucion": a_fecha(r["fecha_ejecucion"]), "comentario": r["comentario"],
        })
    rel = pd.DataFrame(filas, columns=["expediente", "localidad", "departamento", "zona", "desde", "hasta", "informe", "modo",
                                       "km13", "km33", "km66", "estado", "fecha_ejecucion", "comentario"])

    o = obs[obs["expediente"].isin(rel["expediente"])].copy()
    info = rels.set_index("expediente")
    o["localidad"] = o["expediente"].map(info["localidad"])
    o["desde"] = o["expediente"].map(lambda e: a_fecha(info.at[e, "fecha_desde"]))
    o["estado_exp"] = o["expediente"].map(info["estado"])
    o["fecha_ejecucion"] = o["expediente"].map(lambda e: a_fecha(info.at[e, "fecha_ejecucion"]))
    o = o.sort_values(["expediente", "n"]).reset_index(drop=True)
    rel["n_obs"] = rel["expediente"].map(o.groupby("expediente").size()).fillna(0).astype(int)

    todos = [np.asarray(tramos.get(e, []), dtype=np.int64) for e in rel["expediente"]]
    ejec = [s for s, est in zip(todos, rel["estado"]) if est == lg.EJECUTADO]
    pend = [s for s, est in zip(todos, rel["estado"]) if est != lg.EJECUTADO]
    union = lambda L: np.unique(np.concatenate(L)) if L else np.empty(0, dtype=np.int64)  # noqa: E731
    u_all, u_eje, u_pen = union(todos), union(ejec), union(pend)
    n_loc = len(tabla_loc)
    n_loc_rel = int((tabla_loc["estado"] != "Sin relevar").sum())
    km_total = float(red["largo_m"].sum()) / 1000
    km_unico = _km(red, u_all)
    tabla_loc["relevada"] = tabla_loc["estado"] != "Sin relevar"
    res = {
        "fecha": fecha, "n_exp": len(rel), "n_eje": int((rel["estado"] == lg.EJECUTADO).sum()),
        "km_suma": float(rel[["km13", "km33", "km66"]].sum().sum()), "km_unico": km_unico, "km_total": km_total,
        "pct_km": km_unico / km_total * 100 if km_total else 0.0,
        "km_t": {t: _km(red, u_all, t) for t in lg.TENSIONES}, "km_eje": _km(red, u_eje), "km_pend": _km(red, u_pen),
        "n_obs": len(o), "n_loc": n_loc, "n_loc_rel": n_loc_rel, "pct_loc": n_loc_rel / n_loc * 100 if n_loc else 0.0,
        "fuera_de_lista": fuera,
    }
    res["n_pend"] = res["n_exp"] - res["n_eje"]
    return {"rel": rel, "obs": o, "loc": tabla_loc, "res": res}


# ------------------------------------------------------------------ Excel
def excel_bytes(datos: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rel, o, loc, res = datos["rel"], datos["obs"], datos["loc"], datos["res"]
    FUENTE = "Arial"
    f_norm = Font(name=FUENTE, size=10)
    f_neg = Font(name=FUENTE, size=10, bold=True)
    f_tit = Font(name=FUENTE, size=14, bold=True, color="1F3864")
    f_enc = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
    f_azul = Font(name=FUENTE, size=10, color="0000FF")            # valores cargados a mano (convención de las planillas)
    f_link = Font(name=FUENTE, size=10, color="0563C1", underline="single")
    f_nota = Font(name=FUENTE, size=9, italic=True, color="595959")
    rel_enc = PatternFill("solid", start_color="1F3864")
    rel_tot = PatternFill("solid", start_color="DDEBF7")
    fino = Side(style="thin", color="BFBFBF")
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    FMT_FECHA, FMT_KM = "dd/mm/yyyy", "#,##0.0"

    wb = Workbook()
    ws_res = wb.active
    ws_res.title = "Resumen"
    ws_rel = wb.create_sheet("Relevamientos")
    ws_obs = wb.create_sheet("Observaciones")
    ws_loc = wb.create_sheet("Localidades")

    def titulo(ws, texto, sub):
        ws["A1"] = texto
        ws["A1"].font = f_tit
        ws["A2"] = sub
        ws["A2"].font = f_nota

    def encabezados(ws, fila, nombres, anchos):
        for j, (n, w) in enumerate(zip(nombres, anchos), 1):
            c = ws.cell(row=fila, column=j, value=n)
            c.font, c.fill, c.border = f_enc, rel_enc, borde
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.row_dimensions[fila].height = 32

    def celda(ws, fila, col, valor, fmt=None, fuente=None, envolver=False, centrar=False):
        c = ws.cell(row=fila, column=col, value=valor)
        c.font = fuente or f_norm
        c.border = borde
        if fmt:
            c.number_format = fmt
        c.alignment = Alignment(vertical="top", wrap_text=envolver, horizontal="center" if centrar else None)
        return c

    EMISION = f"Emitido el {fecha_corta(res['fecha'].isoformat())}. ENRESP · Gerencia de Energía Eléctrica."
    H0 = 7                                              # primera fila de datos en las hojas de detalle
    n_rel, n_obs, n_loc = len(rel), len(o), len(loc)
    u_rel, u_obs, u_loc = H0 + max(n_rel, 1) - 1, H0 + max(n_obs, 1) - 1, H0 + max(n_loc, 1) - 1

    # ---- Relevamientos (una fila por expediente)
    titulo(ws_rel, "Relevamientos de redes de media tensión", EMISION + " Los totales de la fila 4 se recalculan al filtrar.")
    cols = ["Expediente", "Localidad relevada", "Departamento", "Zona de EDESA", "Relevado desde", "Relevado hasta",
            "Informe", "Cómo se marcó la línea", "Km 13,2 kV", "Km 33 kV", "Km 66 kV", "Km relevados (total)",
            "Observaciones", "Estado", "Fecha de solución por EDESA", "Comentario"]
    encabezados(ws_rel, H0 - 1, cols, [12, 24, 22, 20, 13, 13, 26, 16, 11, 11, 11, 13, 14, 22, 15, 30])
    for i, r in enumerate(rel.itertuples(), H0):
        celda(ws_rel, i, 1, str(r.expediente))
        celda(ws_rel, i, 2, r.localidad, envolver=True)
        celda(ws_rel, i, 3, r.departamento, envolver=True)
        celda(ws_rel, i, 4, r.zona, envolver=True)
        celda(ws_rel, i, 5, r.desde, FMT_FECHA, centrar=True)
        celda(ws_rel, i, 6, r.hasta, FMT_FECHA, centrar=True)
        celda(ws_rel, i, 7, r.informe, envolver=True)
        celda(ws_rel, i, 8, r.modo)
        celda(ws_rel, i, 9, round(r.km13, 3), FMT_KM)
        celda(ws_rel, i, 10, round(r.km33, 3), FMT_KM)
        celda(ws_rel, i, 11, round(r.km66, 3), FMT_KM)
        celda(ws_rel, i, 12, f"=SUM(I{i}:K{i})", FMT_KM, f_neg)
        celda(ws_rel, i, 13, f"=COUNTIF(Observaciones!$A${H0}:$A${u_obs},A{i})", "0", centrar=True)
        celda(ws_rel, i, 14, r.estado)
        celda(ws_rel, i, 15, r.fecha_ejecucion, FMT_FECHA, centrar=True)
        celda(ws_rel, i, 16, r.comentario, envolver=True)
    ws_rel["A4"] = "TOTALES (solo filas visibles)"
    ws_rel["A4"].font = f_neg
    for col in "IJKLM":
        c = ws_rel[f"{col}4"]
        c.value = f"=SUBTOTAL(109,{col}{H0}:{col}{u_rel})"
        c.number_format = "0" if col == "M" else FMT_KM
    for j in range(1, 17):
        ws_rel.cell(row=4, column=j).fill = rel_tot
        ws_rel.cell(row=4, column=j).font = f_neg
        ws_rel.cell(row=4, column=j).border = borde
    ws_rel.freeze_panes = ws_rel.cell(row=H0, column=3)
    ws_rel.auto_filter.ref = f"A{H0 - 1}:P{u_rel}"

    # ---- Observaciones (una fila por observación)
    titulo(ws_obs, "Observaciones de los relevamientos", EMISION + " Una fila por observación del informe.")
    encabezados(ws_obs, H0 - 1, ["Expediente", "Localidad", "N°", "Observación", "Latitud", "Longitud", "Enlace (Google Maps)",
                                 "Fecha de relevamiento", "Estado del expediente", "Fecha de solución por EDESA"],
                [12, 24, 7, 48, 12, 12, 24, 14, 22, 15])
    for i, r in enumerate(o.itertuples(), H0):
        celda(ws_obs, i, 1, str(r.expediente))
        celda(ws_obs, i, 2, r.localidad, envolver=True)
        celda(ws_obs, i, 3, int(r.n), "0", centrar=True)
        celda(ws_obs, i, 4, r.tipo, envolver=True)
        celda(ws_obs, i, 5, None if pd.isna(r.lat) else float(r.lat), "0.000000")
        celda(ws_obs, i, 6, None if pd.isna(r.lon) else float(r.lon), "0.000000")
        c = celda(ws_obs, i, 7, r.enlace or None, fuente=f_link if r.enlace else f_norm)
        if r.enlace:
            c.hyperlink = r.enlace
            c.value = "Abrir en Google Maps"
        celda(ws_obs, i, 8, r.desde, FMT_FECHA, centrar=True)
        celda(ws_obs, i, 9, r.estado_exp)
        celda(ws_obs, i, 10, r.fecha_ejecucion, FMT_FECHA, centrar=True)
    ws_obs["A4"] = "Observaciones visibles"
    ws_obs["C4"] = f"=SUBTOTAL(103,C{H0}:C{u_obs})"
    for j in range(1, 11):
        ws_obs.cell(row=4, column=j).fill = rel_tot
        ws_obs.cell(row=4, column=j).font = f_neg
        ws_obs.cell(row=4, column=j).border = borde
    ws_obs.freeze_panes = ws_obs.cell(row=H0, column=3)
    ws_obs.auto_filter.ref = f"A{H0 - 1}:J{u_obs}"

    # ---- Localidades (la lista de datos/localidades.csv)
    titulo(ws_loc, "Localidades de la provincia y estado de relevamiento",
           EMISION + " Lista: las localidades de la provincia (datos/localidades.csv).")
    encabezados(ws_loc, H0 - 1, ["Localidad", "Departamento", "Población (censo 2010)", "Estado", "Expedientes",
                                 "Fecha de relevamiento", "Fecha de solución por EDESA", "¿Relevada?"],
                [28, 28, 14, 36, 16, 15, 15, 12])
    for i, r in enumerate(loc.itertuples(), H0):
        celda(ws_loc, i, 1, r.nombre)
        celda(ws_loc, i, 2, r.departamento)
        celda(ws_loc, i, 3, None if pd.isna(r.poblacion_2010) or r.poblacion_2010 == 0 else int(r.poblacion_2010), "#,##0")
        celda(ws_loc, i, 4, r.estado)
        celda(ws_loc, i, 5, r.expedientes)
        celda(ws_loc, i, 6, a_fecha(r.fecha_relevamiento), FMT_FECHA, centrar=True)
        celda(ws_loc, i, 7, a_fecha(r.fecha_solucion), FMT_FECHA, centrar=True)
        celda(ws_loc, i, 8, f'=IF(D{i}="Sin relevar","No","Sí")', centrar=True)
    ws_loc["A4"] = "Localidades visibles"
    ws_loc["C4"] = f"=SUBTOTAL(103,A{H0}:A{u_loc})"
    for j in range(1, 9):
        ws_loc.cell(row=4, column=j).fill = rel_tot
        ws_loc.cell(row=4, column=j).font = f_neg
        ws_loc.cell(row=4, column=j).border = borde
    ws_loc.freeze_panes = ws_loc.cell(row=H0, column=2)
    ws_loc.auto_filter.ref = f"A{H0 - 1}:H{u_loc}"

    # ---- Resumen (fórmulas sobre las otras hojas)
    titulo(ws_res, "Resumen de relevamientos de redes de media tensión", EMISION)
    ws_res["A3"] = "Si las celdas aparecen vacías o en cero, hacé clic en «Habilitar edición»: Excel recalcula las fórmulas al abrir."
    ws_res["A3"].font = f_nota
    for j, (n, w) in enumerate(zip(["Indicador", "Valor", "Cómo se calcula / fuente"], [58, 16, 90]), 1):
        c = ws_res.cell(row=5, column=j, value=n)
        c.font, c.fill, c.border = f_enc, rel_enc, borde
        c.alignment = Alignment(horizontal="center")
        ws_res.column_dimensions[get_column_letter(j)].width = w
    R = f"Relevamientos!"
    indic = [
        ("Expedientes relevados", f"=COUNTA({R}A{H0}:A{u_rel})" if n_rel else 0, "0", "Una fila por expediente en la hoja Relevamientos.", False),
        ("  con adecuaciones ejecutadas por EDESA", f'=COUNTIF({R}N{H0}:N{u_rel},"{lg.EJECUTADO}")', "0", "Expedientes en estado «Ejecutado por EDESA».", False),
        ("  pendientes de ejecución por EDESA", "=B6-B7", "0", "Relevados menos ejecutados.", False),
        ("Km de red MT relevados (suma de los expedientes)", f"=SUM({R}L{H0}:L{u_rel})", FMT_KM, "Suma de la hoja Relevamientos. Si dos expedientes comparten tramos, esos tramos se suman dos veces.", False),
        ("Km de red MT relevados (sin repetir tramos comunes)", round(res["km_unico"], 3), FMT_KM, "Calculado por la aplicación uniendo los tramos de todos los expedientes. Es el valor que se usa para el porcentaje.", True),
        ("Km de red MT de la provincia", round(res["km_total"], 3), FMT_KM, FUENTE_RED + ".", True),
        ("% de la red MT relevada", "=IF(B11>0,B10/B11,0)", "0.0%", "Km relevados sin repetir ÷ km de red de la provincia.", False),
        ("  Km relevados en 13,2 kV", f"=SUM({R}I{H0}:I{u_rel})", FMT_KM, "Suma de la hoja Relevamientos.", False),
        ("  Km relevados en 33 kV", f"=SUM({R}J{H0}:J{u_rel})", FMT_KM, "Suma de la hoja Relevamientos.", False),
        ("  Km relevados en 66 kV", f"=SUM({R}K{H0}:K{u_rel})", FMT_KM, "Suma de la hoja Relevamientos.", False),
        ("Observaciones registradas", f"=COUNTA(Observaciones!A{H0}:A{u_obs})" if n_obs else 0, "0", "Filas de la hoja Observaciones.", False),
        ("Localidades de la provincia", f"=COUNTA(Localidades!A{H0}:A{u_loc})" if n_loc else 0, "0", "Las localidades de la lista de la provincia (datos/localidades.csv).", False),
        ("Localidades relevadas", f'=COUNTIF(Localidades!H{H0}:H{u_loc},"Sí")', "0", "Localidades que figuran en la columna «Localidad» de algún expediente.", False),
        ("% de localidades relevadas", "=IF(B17>0,B18/B17,0)", "0.0%", "Localidades relevadas ÷ localidades de la provincia.", False),
        ("Localidades pendientes de relevamiento", "=B17-B18", "0", "Detalle en la hoja Localidades (filtrar «¿Relevada?» = No).", False),
    ]
    for k, (nombre, valor, fmt, nota, manual) in enumerate(indic, 6):
        celda(ws_res, k, 1, nombre, fuente=f_neg if not nombre.startswith("  ") else f_norm)
        celda(ws_res, k, 2, valor, fmt, f_azul if manual else f_neg)
        celda(ws_res, k, 3, nota, envolver=True, fuente=f_nota)
    ws_res["B10"].comment = Comment("Valor cargado por la aplicación al generar el reporte (no es una fórmula de la planilla).", "ENRESP")
    ws_res["B11"].comment = Comment(FUENTE_RED, "ENRESP")
    ws_res["A22"] = "Azul: valor cargado por la aplicación. Negro: fórmula de la planilla."
    ws_res["A22"].font = f_nota
    ws_res.sheet_view.showGridLines = False
    if res["fuera_de_lista"]:
        ws_res["A23"] = ("Localidades escritas en expedientes que no están en la lista de la provincia (no cuentan en el porcentaje): "
                         + ", ".join(res["fuera_de_lista"]) + ".")
        ws_res["A23"].font = f_nota

    for ws in (ws_rel, ws_obs, ws_loc):
        ws.sheet_view.zoomScale = 100
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = f"{H0 - 1}:{H0 - 1}"
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------- Word
def _fuente_run(run, negrita=False, tam=11):
    from docx.oxml.ns import qn
    from docx.shared import Pt
    run.font.name = "Arial"
    run.font.size = Pt(tam)
    run.bold = negrita
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for att in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(att), "Arial")


def _base_word(plantilla):
    """Documento con el formato de la nota modelo y las funciones para escribir párrafos y tablas."""
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    if plantilla and Path(plantilla).exists():
        doc = Document(str(plantilla))              # se conserva el formato de página de la nota modelo
        cuerpo = doc.element.body
        for hijo in list(cuerpo):
            if hijo.tag != qn("w:sectPr"):
                cuerpo.remove(hijo)
    else:
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = s.top_margin = Cm(3.0)
        s.bottom_margin = Cm(2.5)

    def par(partes, alinear=None, sangria=False, interlineado=None, despues=None, juntar=False):
        p = doc.add_paragraph()
        if isinstance(partes, str):
            partes = [(partes, False)]
        for texto, neg in partes:
            _fuente_run(p.add_run(texto), neg)
        if alinear is not None:
            p.alignment = alinear
        if sangria:
            p.paragraph_format.first_line_indent = Cm(2.5)
        if interlineado:
            p.paragraph_format.line_spacing = interlineado
        if despues is not None:
            p.paragraph_format.space_after = Pt(despues)
        if juntar:
            p.paragraph_format.keep_with_next = True
        return p

    def borde_tabla(tabla):
        tbl_pr = tabla._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{lado}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "808080")
            b.append(e)
        tbl_pr.append(b)

    def tabla(encabezado, filas, anchos_cm, derecha=()):
        t = doc.add_table(rows=1, cols=len(encabezado))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        borde_tabla(t)
        todas = [encabezado] + filas
        for i, fila in enumerate(todas):
            celdas = t.rows[0].cells if i == 0 else t.add_row().cells
            for j, (txt, w) in enumerate(zip(fila, anchos_cm)):
                c = celdas[j]
                c.width = Cm(w)
                p = c.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                if j in derecha and i > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    sh = OxmlElement("w:shd")
                    sh.set(qn("w:val"), "clear")
                    sh.set(qn("w:color"), "auto")
                    sh.set(qn("w:fill"), "E7E6E6")
                    c._tc.get_or_add_tcPr().append(sh)
                _fuente_run(p.add_run(str(txt)), negrita=(i == 0), tam=9)
        pr = t.rows[0]._tr.get_or_add_trPr()
        h = OxmlElement("w:tblHeader")
        h.set(qn("w:val"), "true")
        pr.append(h)
        for fila_ in t.rows:                                    # que una fila no se parta entre dos páginas
            cs = OxmlElement("w:cantSplit")
            cs.set(qn("w:val"), "true")
            fila_._tr.get_or_add_trPr().append(cs)
        return t

    return doc, par, tabla


def word_bytes(datos: dict, fecha: date, referencia: str, plantilla: Path | None = None) -> bytes:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    rel, loc, res = datos["rel"], datos["loc"], datos["res"]
    doc, par, tabla = _base_word(plantilla)
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    par(f"Salta, {fecha_larga(fecha)}", alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    par([("A: Directorio", True)])
    par([("De: Gerencia Eléctrica", True)])
    par("", alinear=J, interlineado=1.5)
    par([("Ref:", True), (" " + referencia, False)], alinear=J, interlineado=1.5)
    par("")

    km_t = res["km_t"]
    par("Por medio de la presente se informa el estado de avance de los relevamientos de campo de las redes de media tensión "
        "(13,2, 33 y 66 kV) de la Provincia de Salta que realiza esta Gerencia, y el seguimiento de las adecuaciones "
        f"solicitadas a la Distribuidora EDESA S.A., al {fecha_larga(fecha)}.", alinear=J, sangria=True, interlineado=1.5)
    par(f"A la fecha se relevaron {res['n_loc_rel']} de los {res['n_loc']} municipios de la Provincia, "
        f"lo que representa el {num(res['pct_loc'])} %, y {num(res['km_unico'])} km de los {num(res['km_total'])} km de red de "
        f"media tensión registrados, equivalentes al {num(res['pct_km'])} % de la red "
        f"({num(km_t[13.2])} km en 13,2 kV, {num(km_t[33.0])} km en 33 kV y {num(km_t[66.0])} km en 66 kV).",
        alinear=J, sangria=True, interlineado=1.5)
    if res["n_exp"]:
        par(f"En los {res['n_exp']} expedientes de relevamiento se registraron {res['n_obs']} observaciones. "
            f"En {res['n_eje']} de ellos EDESA ya ejecutó las adecuaciones solicitadas ({num(res['km_eje'])} km de red), "
            f"mientras que {res['n_pend']} permanecen pendientes de ejecución ({num(res['km_pend'])} km de red).",
            alinear=J, sangria=True, interlineado=1.5)
        par("El detalle por expediente es el siguiente:", alinear=J, sangria=True, interlineado=1.5, juntar=True)

    if res["n_exp"]:
        filas = []
        for r in rel.itertuples():
            fechas = fecha_corta(r.desde.isoformat()) if r.desde else ""
            if r.hasta and r.hasta != r.desde:
                fechas += f" al {fecha_corta(r.hasta.isoformat())}"
            if r.estado == lg.EJECUTADO:
                est = "Adecuaciones ejecutadas por EDESA" + (f" el {fecha_corta(r.fecha_ejecucion.isoformat())}" if r.fecha_ejecucion else "")
            else:
                est = "Relevado, adecuaciones pendientes"
            filas.append([r.expediente, r.localidad, fechas, num(r.km13 + r.km33 + r.km66), r.n_obs, est])
        tabla(["Expediente", "Municipio", "Relevamiento", "Km de red", "Obser-vaciones", "Estado"], filas,
              [2.0, 3.2, 3.0, 1.7, 1.9, 3.2], derecha=(3, 4))
        par("", despues=4)

    pend = loc[~loc["relevada"]]
    if len(pend):
        par(f"Los {len(pend)} municipios pendientes de relevamiento, agrupadas por departamento, son:",
            alinear=J, sangria=True, interlineado=1.5, juntar=True)
        filas = [[d, ", ".join(g["nombre"].sort_values())] for d, g in pend.groupby("departamento")]
        filas.sort(key=lambda f: f[0])
        tabla(["Departamento", "Municipios pendientes"], filas, [4.2, 10.8])
        par("", despues=4)
    else:
        par("No quedan municipios pendientes de relevamiento.", alinear=J, sangria=True, interlineado=1.5)

    par("Atentamente.", alinear=J, sangria=True, interlineado=1.5)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def word_localidades_bytes(datos: dict, nombres: list[str], fecha: date, referencia: str,
                           plantilla: Path | None = None) -> bytes:
    """Nota en Word (mismo formato que la nota general) sobre las localidades elegidas."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    rel, loc, res = datos["rel"], datos["loc"], datos["res"]
    doc, par, tabla = _base_word(plantilla)
    J = WD_ALIGN_PARAGRAPH.JUSTIFY
    sel = loc[loc["nombre"].isin(nombres)].sort_values(["departamento", "nombre"])
    n = len(sel)
    sin = sel[sel["estado"] == "Sin relevar"]
    eje = sel[sel["estado"] == "Adecuaciones ejecutadas por EDESA"]
    pen = sel[sel["estado"] == "Relevada · adecuaciones pendientes"]
    exps = sorted({e.strip() for v in sel["expedientes"] for e in str(v).split(",") if e.strip()})
    rel_s = rel[rel["expediente"].isin(exps)]
    km_s = float(rel_s[["km13", "km33", "km66"]].sum().sum())
    n_obs = int(rel_s["n_obs"].sum()) if len(rel_s) else 0

    par(f"Salta, {fecha_larga(fecha)}", alinear=WD_ALIGN_PARAGRAPH.RIGHT)
    par([("A: Directorio", True)])
    par([("De: Gerencia Eléctrica", True)])
    par("", alinear=J, interlineado=1.5)
    par([("Ref:", True), (" " + referencia, False)], alinear=J, interlineado=1.5)
    par("")
    par("Por medio de la presente se informa el estado de los relevamientos de campo de las redes de media tensión "
        f"(13,2, 33 y 66 kV) en los {n} municipios indicados de la Provincia de Salta y el seguimiento de las "
        f"adecuaciones solicitadas a la Distribuidora EDESA S.A., al {fecha_larga(fecha)}.",
        alinear=J, sangria=True, interlineado=1.5)
    def pl(k, uno, varios):
        return f"{k} {uno if k == 1 else varios}"
    nr = n - len(sin)
    texto = (f"{'Del municipio considerado' if n == 1 else f'De los {n} municipios considerados'}, "
             f"{'fue relevado' if n == 1 and nr else ('no fue relevado' if n == 1 else f'{nr} fueron relevados')}"
             f"{'' if n == 1 else f' y {len(sin)} están pendientes de relevamiento'}.")
    if nr:
        texto += (f" Entre los relevados, en {len(eje)} EDESA ya ejecutó las adecuaciones solicitadas y "
                  f"en {len(pen)} permanecen pendientes de ejecución.")
    if len(rel_s):
        texto += (f" {'El expediente que los incluye suma' if len(rel_s) == 1 else f'Los {len(rel_s)} expedientes que los incluyen suman'} "
                  f"{num(km_s)} km de red y {pl(n_obs, 'observación', 'observaciones')} "
                  "(un expediente puede abarcar otros municipios además de los consultados).")
    par(texto, alinear=J, sangria=True, interlineado=1.5)
    if n:
        par("El estado por municipio es el siguiente:", alinear=J, sangria=True, interlineado=1.5, juntar=True)
        filas = [[r.nombre, r.departamento, r.estado.replace("Relevada · a", "Relevada, a"), r.expedientes or "-",
                  fecha_corta(r.fecha_relevamiento) or "-", fecha_corta(r.fecha_solucion) or "-"] for r in sel.itertuples()]
        tabla(["Municipio", "Departamento", "Estado", "Expediente", "Relevamiento", "Solución EDESA"], filas,
              [2.7, 2.9, 3.3, 2.0, 2.0, 2.1])
        par("", despues=4)
    if len(rel_s):
        par("El detalle de los expedientes correspondientes es el siguiente:", alinear=J, sangria=True, interlineado=1.5, juntar=True)
        filas = []
        for r in rel_s.itertuples():
            fechas = fecha_corta(r.desde.isoformat()) if r.desde else ""
            if r.hasta and r.hasta != r.desde:
                fechas += f" al {fecha_corta(r.hasta.isoformat())}"
            if r.estado == lg.EJECUTADO:
                est = "Adecuaciones ejecutadas por EDESA" + (f" el {fecha_corta(r.fecha_ejecucion.isoformat())}" if r.fecha_ejecucion else "")
            else:
                est = "Relevado, adecuaciones pendientes"
            filas.append([r.expediente, r.localidad, fechas, num(r.km13 + r.km33 + r.km66), r.n_obs, est])
        tabla(["Expediente", "Municipio", "Relevamiento", "Km de red", "Obser-vaciones", "Estado"], filas,
              [2.0, 3.2, 3.0, 1.7, 1.9, 3.2], derecha=(3, 4))
        par("", despues=4)
    if len(sin):
        par(f"Los {len(sin)} municipios pendientes de relevamiento, agrupadas por departamento, son:" if len(sin) > 1
            else "El municipio pendiente de relevamiento es:",
            alinear=J, sangria=True, interlineado=1.5, juntar=True)
        filas = [[d, ", ".join(g["nombre"].sort_values())] for d, g in sin.groupby("departamento")]
        filas.sort(key=lambda f: f[0])
        tabla(["Departamento", "Municipios pendientes"], filas, [4.2, 10.8])
        par("", despues=4)
    elif n:
        par("Todos los municipios considerados fueron relevados.", alinear=J, sangria=True, interlineado=1.5)
    par("Las cifras de red surgen de la extracción regulatoria de EDESA (red de media tensión de 13,2, 33 y 66 kV).",
        alinear=J, sangria=True, interlineado=1.5)
    par("Atentamente.", alinear=J, sangria=True, interlineado=1.5)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
